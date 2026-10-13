#!/usr/bin/env python3
"""Apply a returned review workbook to the evalset.

    python review/apply_review.py review/returned/review_ar_sy.xlsx            # dry run: show what would change
    python review/apply_review.py review/returned/review_ar_sy.xlsx --write    # update evalset/ar_sy.jsonl

Questions sheet, per row (matched by ID):
    Correction filled  -> question replaced; the old text is kept in question_before_review
    OK? = ✓ or a Correction -> translation_reviewed = true
    Comment            -> kept in review_comment
Terms sheet: corrections are listed, not applied (a term appears in many questions; change those by hand,
then rebuild the workbooks with make_review.py).
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

EVALSET = Path(__file__).resolve().parent.parent / "evalset"


def cell_text(v) -> str:
    return str(v).strip() if v is not None else ""


def read_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    head = [cell_text(h) for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(v is not None for v in r)]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--set", help="evalset file name, e.g. ar_sy (default: taken from review_<set>.xlsx)")
    ap.add_argument("--write", action="store_true", help="write the changes (default: dry run)")
    args = ap.parse_args()
    name = args.set or re.sub(r"^review_", "", args.workbook.stem)
    path = EVALSET / f"{name}.jsonl"
    if not path.exists():
        raise SystemExit(f"no evalset file {path}; pass --set")

    wb = load_workbook(args.workbook)
    items = [json.loads(l) for l in path.open(encoding="utf-8") if l.strip()]
    by_id = {r["id"]: r for r in items}
    changed, confirmed, unknown = [], [], []
    for row in read_sheet(wb["Questions"]):
        qid = cell_text(row.get("ID"))
        if qid not in by_id:
            unknown.append(qid)
            continue
        item = by_id[qid]
        correction, ok, comment = (cell_text(row.get(k)) for k in ("Correction", "OK?", "Comment"))
        if correction and correction != item["question"]:
            changed.append((qid, item["question"], correction))
            item.setdefault("question_before_review", item["question"])
            item["question"] = correction
        if correction or ok:
            item["translation_reviewed"] = True
            confirmed.append(qid)
        if comment:
            item["review_comment"] = comment

    print(f"{path.name}: {len(confirmed)} reviewed, {len(changed)} corrected, "
          f"{len(items) - len(confirmed)} not marked")
    for qid, old, new in changed:
        print(f"  {qid}: {old}\n     -> {new}")
    if unknown:
        print("  IDs not in the evalset (ignored):", ", ".join(unknown))
    if "Terms" in wb.sheetnames:
        terms = [(cell_text(r.get("Term in the questions")), cell_text(r.get("Correction")), cell_text(r.get("Comment")))
                 for r in read_sheet(wb["Terms"])]
        todo = [t for t in terms if t[1] or t[2]]
        if todo:
            print("\nTerm corrections (apply by hand, they affect many questions):")
            for term, corr, com in todo:
                print(f"  {term} -> {corr or '(no correction)'}" + (f"   # {com}" if com else ""))
    if args.write:
        path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in items), encoding="utf-8")
        print(f"\nwritten: {path}")
    else:
        print("\ndry run: nothing written (add --write)")


if __name__ == "__main__":
    main()
