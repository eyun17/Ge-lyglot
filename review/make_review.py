#!/usr/bin/env python3
"""Build the native-speaker review workbooks (one .xlsx per language set) from evalset/.

    python review/make_review.py          # writes review/review_<set>.xlsx
    python review/apply_review.py ...     # applies a returned workbook to evalset/ (see that file)

Each workbook has three sheets (German: no Terms sheet, the terms are the German originals):
    Instructions   what to do, in English and in the reviewer's language
    Questions      one row per question: English, current translation, OK?, Correction, Comment
    Terms          recurring words for German offices and permits, to check once
Korean and English are maintained by the project owner and get no workbook.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
EVALSET = ROOT / "evalset"
OUT = Path(__file__).resolve().parent
OK_MARK = "✓"

AR_TERMS = [("البطاقة الزرقاء", "Blaue Karte EU"), ("أوسبيلدونغ", "Ausbildung"),
            ("الإقامة الدائمة", "Niederlassungserlaubnis"), ("دائرة الأجانب", "Ausländerbehörde"),
            ("الوكالة الاتحادية للعمل", "Bundesagentur für Arbeit"), ("بطاقة الفرص", "Chancenkarte"),
            ("أوبير", "Au-pair"), ("تسجيل السكن", "Anmeldung"), ("إقامة البحث عن شغل / عمل", "Aufenthaltserlaubnis zur Arbeitsplatzsuche"),
            ("Working Holiday", "Working-Holiday-Visum"), ("التأمين الاجتماعي للفنانين", "Künstlersozialversicherung"),
            ("Werkstudent / HiWi", "Werkstudent / HiWi"), ("الميني جوب", "Minijob"),
            ("الكيندرغيلد", "Kindergeld"), ("الرقم الضريبي", "Steuer-ID"), ("رسوم الـ GEZ", "Rundfunkbeitrag"),
            ("إلغاء تسجيل السكن", "Abmeldung"), ("إعانة البطالة", "Arbeitslosengeld")]
IT_TERMS = [("Carta Blu UE", "Blaue Karte EU"), ("fare un'Ausbildung", "Ausbildung"),
            ("permesso di soggiorno permanente", "Niederlassungserlaubnis"), ("ufficio stranieri", "Ausländerbehörde"),
            ("Agenzia federale per il lavoro", "Bundesagentur für Arbeit"), ("Carta delle opportunità", "Chancenkarte"),
            ("au pair", "Au-pair"), ("fare l'Anmeldung", "Anmeldung"),
            ("permesso di soggiorno di 18 mesi per cercare lavoro", "Aufenthaltserlaubnis zur Arbeitsplatzsuche"),
            ("visto working holiday", "Working-Holiday-Visum"), ("previdenza sociale per artisti", "Künstlersozialversicherung"),
            ("Werkstudent / HiWi", "Werkstudent / HiWi"), ("Minijob", "Minijob"),
            ("Kindergeld", "Kindergeld"), ("Steuer-ID", "Steuer-ID"), ("Rundfunkbeitrag", "Rundfunkbeitrag"),
            ("fare l'Abmeldung", "Abmeldung"), ("indennità di disoccupazione", "Arbeitslosengeld")]

SETS = {
    "ar": {
        "title": "Arabic (Modern Standard Arabic)", "column": "Arabic", "rtl": True, "terms": AR_TERMS,
        "extra": ["This set is written in Modern Standard Arabic on purpose: it is the baseline. Please keep it "
                  "standard Arabic. A separate file covers the Syrian dialect.",
                  "Some questions use the masculine form (for example مواطن). If a neutral or feminine form would be "
                  "more natural, say so in Comment."],
        "ai_note": ("هذه الجمل مسودة أولى كتبها نموذج ذكاء اصطناعي (Claude Opus 5.5)، وقد تحتوي على أخطاء أو "
                    "صياغات غير طبيعية، فلا تتردد في تصحيحها كما تكتبها أنت."),
        "summary": ("ملخص: في ورقة Questions، اقرأ الجملة الإنجليزية ثم الترجمة العربية. إذا كانت صحيحة وطبيعية، "
                    "اختر ✓ في عمود OK? ولا تكتب شيئاً آخر. إذا كان فيها خطأ أو كانت غير طبيعية، اكتب صيغتك في عمود "
                    "Correction. لا تغيّر المعنى، صحّح اللغة فقط، وأبقِ الأسئلة بالفصحى. شكراً لك!"),
    },
    "ar_sy": {
        "title": "Syrian Arabic (Levantine dialect)", "column": "Syrian Arabic", "rtl": True, "terms": AR_TERMS,
        "msa": True,
        "extra": ["These questions are written in Syrian (Levantine) colloquial Arabic on purpose, the way people type "
                  "in a chat. Correct them to how you would actually write it; any Syrian city's dialect is fine.",
                  "Mixing in German words (Ausbildung, Duldung) is fine if people really do that.",
                  "The column 'MSA (reference only)' shows the standard Arabic version; you don't need to review it."],
        "ai_note": ("هالجمل مسودة أولى كتبها نموذج ذكاء اصطناعي (Claude Opus 5.5)، وممكن يكون فيها غلطات أو حكي "
                    "مو طبيعي، فصحّحها براحتك متل ما بتكتبها إنت."),
        "summary": ("ملخص: هالأسئلة مكتوبة باللهجة السورية عن قصد، متل ما الناس بتكتب بالشات. إذا السؤال طبيعي، "
                    "اختار ✓ بعمود OK?. إذا مو طبيعي، اكتب كيف كنت رح تكتبه بعمود Correction. "
                    "ما تغيّر المعنى. عمود الفصحى للمرجع بس. شكراً!"),
    },
    "it": {
        "title": "Italian", "column": "Italian", "rtl": False, "terms": IT_TERMS,
        "extra": ["Some questions use the masculine form (for example \"Sono stato in Germania\", \"sono siriano\"). "
                  "If a neutral or feminine form would be more natural, say so in Comment."],
        "ai_note": ("Queste frasi sono una prima bozza scritta da un modello di intelligenza artificiale (Claude Opus 5.5): "
                    "possono contenere errori o suonare poco naturali, quindi correggile liberamente, come le scriveresti tu."),
        "summary": ("Riassunto: nel foglio Questions leggi la frase in inglese e la traduzione italiana. Se è corretta e "
                    "suona naturale, scegli ✓ nella colonna OK? e non scrivere altro. Se c'è un errore o non suona "
                    "naturale, scrivi la tua versione nella colonna Correction. Non cambiare il significato, correggi "
                    "solo la lingua. Grazie!"),
    },
    "de": {
        "title": "German", "column": "German", "rtl": False, "terms": [],
        "extra": ["German questions should sound like a non-lawyer typing into a chat: everyday words are fine "
                  "(\"jobben\", \"Visum wechseln\"), legal precision is not needed.",
                  "Some questions use the masculine form (for example \"Ich bin Syrer\"). If a neutral or feminine "
                  "form would be more natural, say so in Comment."],
        "ai_note": ("Diese Sätze sind ein erster Entwurf, den ein KI-Modell (Claude Opus 5.5) geschrieben hat. Sie können "
                    "Fehler enthalten oder unnatürlich klingen, also korrigiere sie gern so, wie du sie selbst schreiben würdest."),
        "summary": ("Kurz gesagt: Lies im Blatt Questions den englischen Satz und dann die deutsche Übersetzung. Wenn sie "
                    "richtig ist und natürlich klingt, wähle ✓ in der Spalte OK? und schreibe sonst nichts. Wenn etwas falsch "
                    "ist oder unnatürlich klingt, schreibe deine Version in die Spalte Correction. Ändere nicht den Inhalt, "
                    "nur die Sprache. Danke!"),
    },
}

HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")
WRAP = Alignment(wrap_text=True, vertical="top")
RTL = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)


def load(name: str) -> list[dict]:
    return [json.loads(l) for l in (EVALSET / f"{name}.jsonl").open(encoding="utf-8") if l.strip()]


def instructions(ws, cfg: dict, n: int) -> None:
    lines = [
        (f"Translation review — {cfg['title']}", Font(bold=True, size=14)),
        ("", None),
        ("GesetzPolyglot answers questions about German law (residence, work, tax, social insurance) in the "
         "user's own language. The questions in this file are used to test it, so they should read like real "
         "questions that someone from your community would type into a chat.", None),
        ("", None),
        ("The same 60 questions are asked in every language. Many name the asker's nationality (American, Syrian, "
         "Portuguese ...); this is part of the test, so please keep the nationality even if it is not yours.", None),
        ("", None),
        ("These sentences are a first draft written by an AI model (Claude Opus 5.5). They may sound unnatural or "
         "contain mistakes, so please correct them freely, as you would write them yourself.", Font(bold=True)),
        ("", None),
        ("How to review (about 30–45 minutes)", Font(bold=True)),
        (f"1. Go to the sheet 'Questions'. There are {n} questions, one per row.", None),
        (f"2. Read the English sentence, then the {cfg['column']} sentence next to it.", None),
        (f"3. If it is correct and sounds natural: choose {OK_MARK} in the column 'OK?'. Nothing else to do.", None),
        ("4. If something is wrong or unnatural: write your version in 'Correction'. Keep the meaning, fix only the language.", None),
        ("5. Unsure? Write a short note in 'Comment'.", None),
        *([("6. Then check the sheet 'Terms' once: words for German offices and permits that appear in many questions.", None)]
          if cfg["terms"] else []),
        ("", None),
        ("Please check especially", Font(bold=True)),
        ("• Words for German offices and permits: use the words people in your community actually use. "
         "Please do not add the German word in parentheses.", None),
        ("• Rows marked 🌍 name the asker's nationality. Check that this sounds right.", None),
        ("• Everyday wording is better than formal or legal wording.", None),
        *[("• " + e, None) for e in cfg["extra"]],
        ("", None),
        ("Please don't change the topic of a question or the numbers in it (for example 48,000 euros, 32 years old).", None),
        ("When you are done, send this file back. Thank you!", Font(bold=True)),
        ("", None),
        (cfg["ai_note"], Font(bold=True)),
        (cfg["summary"], None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = RTL if (cfg["rtl"] and text in (cfg["summary"], cfg["ai_note"])) else WRAP
        if font:
            c.font = font
    ws.column_dimensions["A"].width = 110


def header(ws, cols: list[tuple[str, int]]) -> None:
    for j, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill, c.alignment = Font(bold=True), HEADER_FILL, WRAP
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"


def ok_dropdown(ws, col_letter: str, last_row: int) -> None:
    dv = DataValidation(type="list", formula1=f'"{OK_MARK}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")


def build(name: str, cfg: dict) -> Path:
    rows, en = load(name), {r["id"]: r["question"] for r in load("en")}
    wb = Workbook()
    instructions(wb.active, cfg, len(rows))
    wb.active.title = "Instructions"

    ws = wb.create_sheet("Questions")
    cols = [("No", 5), ("ID", 7), ("English", 48), (cfg["column"], 48)]
    if cfg.get("msa"):
        cols.append(("MSA (reference only)", 40))
    cols += [("OK?", 6), ("Correction", 48), ("Comment", 30)]
    header(ws, cols)
    names = [c[0] for c in cols]
    for i, r in enumerate(rows, 2):
        english = r.get("question_en") or en[r["id"]]
        values = {"No": i - 1, "ID": r["id"], "English": ("🌍 " if r.get("persona_country") else "") + english,
                  cfg["column"]: r["question"], "MSA (reference only)": r.get("question_msa", "")}
        for j, col in enumerate(names, 1):
            c = ws.cell(row=i, column=j, value=values.get(col))
            c.alignment = RTL if (cfg["rtl"] and col in (cfg["column"], "MSA (reference only)", "Correction")) else WRAP
    ok_dropdown(ws, ws.cell(row=1, column=names.index("OK?") + 1).column_letter, len(rows) + 1)

    if cfg["terms"]:
        wt = wb.create_sheet("Terms")
        header(wt, [("Term in the questions", 32), ("German original (for you)", 28), ("OK?", 6),
                    ("Correction", 32), ("Comment", 30)])
        for i, (term, de) in enumerate(cfg["terms"], 2):
            wt.cell(row=i, column=1, value=term).alignment = RTL if cfg["rtl"] else WRAP
            wt.cell(row=i, column=2, value=de).alignment = WRAP
        ok_dropdown(wt, "C", len(cfg["terms"]) + 1)

    wb.active = 1  # open on the questions
    path = OUT / f"review_{name}.xlsx"
    wb.save(path)
    return path


def main() -> None:
    for name, cfg in SETS.items():
        path = build(name, cfg)
        print(f"{path.relative_to(ROOT)}: {len(load(name))} questions, {len(cfg['terms'])} terms")


if __name__ == "__main__":
    main()
